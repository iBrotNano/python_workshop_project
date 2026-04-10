#!/usr/bin/env python3
import csv
import unicodedata
import re
import logging

from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook
from common.progress_callback import ProgressCallback, emit_progress

log = logging.getLogger(__name__)


class SwissNutritionDbXlsxToCsvConverter:
    """
    Converts the Swiss nutrition database from its original Excel format to a cleaned CSV format.
    The converter maps the original German column headers to Open Food Facts compatible English names,
    normalizes numeric values to grams per 100g where necessary, and handles missing or non-applicable data.
    """

    COLUMN_MAP = {
        "id": "code",
        "name": "product_name",
        "synonyme": "synonyms_text",
        "synonyma": "synonyms_text",  # guards against alternate headers
        "synonyms": "synonyms_text",
        "kategorie": "categories_text",
        "energie_kilojoule_kj": "energy_100g_kj",
        "energie_kilojoule": "energy_100g_kj",
        "fett_total_g": "fat_100g",
        "fett_total": "fat_100g",
        "fettsaeuren_gesaettigt_g": "saturated_fat_100g",
        "fettsaeuren_einfach_ungesaettigt_g": "monounsaturated_fat_100g",
        "fettsaeuren_mehrfach_ungesaettigt_g": "polyunsaturated_fat_100g",
        "linolsaeure_g": "linoleic_acid_100g",
        "alpha_linolensaeur_g": "alpha_linolenic_acid_100g",
        "alpha_linolensaure_g": "alpha_linolenic_acid_100g",
        "alpha_linolensaeure_g": "alpha_linolenic_acid_100g",
        "alpha_linolenic_acid_g": "alpha_linolenic_acid_100g",
        "eicosapentaensaure_epa_g": "eicosapentaenoic_acid_100g",
        "eicosapentaensaeure_epa_g": "eicosapentaenoic_acid_100g",
        "docosahexaensaure_dha_g": "docosahexaenoic_acid_100g",
        "docosahexaensaeure_dha_g": "docosahexaenoic_acid_100g",
        "cholesterin_mg": "cholesterol_100g",
        "kohlenhydrate_verfugbar_g": "carbohydrates_100g",
        "kohlenhydrate_verfuegbar_g": "carbohydrates_100g",
        "zucker_g": "sugars_100g",
        "starke_g": "starch_100g",
        "starch_g": "starch_100g",
        "staerke_g": "starch_100g",
        "nahrungsfasern_g": "fiber_100g",
        "protein_g": "proteins_100g",
        "salz_nacl_g": "salt_100g",
        "alkohol_g": "alcohol_100g",
        "retinol_ug": "vitamin_a_100g",
        "retinol_g": "vitamin_a_100g",
        "vitamin_a_aktivitat_re_g": "vitamin_a_100g",
        "vitamin_a_aktivitat_rae_g": "vitamin_a_100g",
        "betacarotin_ug": "beta_carotene_100g",
        "betacarotin_g": "beta_carotene_100g",
        "betacarotin_aktivitat_g": "beta_carotene_100g",
        "vitamin_b1_thiamin_mg": "vitamin_b1_100g",
        "vitamin_b2_riboflavin_mg": "vitamin_b2_100g",
        "niacin_mg": "vitamin_b3_100g",
        "vitamin_b6_pyridoxin_mg": "vitamin_b6_100g",
        "folat_ug": "vitamin_b9_100g",
        "folat_g": "vitamin_b9_100g",
        "vitamin_b12_cobalamin_ug": "vitamin_b12_100g",
        "vitamin_b12_cobalamin_g": "vitamin_b12_100g",
        "pantothensaure_mg": "pantothenic_acid_100g",
        "pantothensaeure_mg": "pantothenic_acid_100g",
        "vitamin_c_ascorbinsaure_mg": "vitamin_c_100g",
        "vitamin_c_ascorbinsaeure_mg": "vitamin_c_100g",
        "vitamin_d_calciferol_ug": "vitamin_d_100g",
        "vitamin_d_calciferol_g": "vitamin_d_100g",
        "vitamin_e_a_tocopherol_mg": "vitamin_e_100g",
        "vitamin_e_tocopherol_mg": "vitamin_e_100g",
        "kalium_k_mg": "potassium_100g",
        "natrium_na_mg": "sodium_100g",
        "chlorid_cl_mg": "chloride_100g",
        "calcium_ca_mg": "calcium_100g",
        "magnesium_mg_mg": "magnesium_100g",
        "magnesium_mg": "magnesium_100g",
        "phosphor_p_mg": "phosphorus_100g",
        "eisen_fe_mg": "iron_100g",
        "jod_i_ug": "iodine_100g",
        "jod_i_g": "iodine_100g",
        "zink_zn_mg": "zinc_100g",
        "selen_se_ug": "selenium_100g",
        "selen_se_g": "selenium_100g",
    }

    OUTPUT_COLUMNS = [
        "code",
        "product_name",
        "synonyms_text",
        "categories_text",
        "energy_100g_kj",
        "fat_100g",
        "saturated_fat_100g",
        "monounsaturated_fat_100g",
        "polyunsaturated_fat_100g",
        "linoleic_acid_100g",
        "alpha_linolenic_acid_100g",
        "eicosapentaenoic_acid_100g",
        "docosahexaenoic_acid_100g",
        "cholesterol_100g",
        "carbohydrates_100g",
        "sugars_100g",
        "starch_100g",
        "fiber_100g",
        "proteins_100g",
        "salt_100g",
        "alcohol_100g",
        "vitamin_a_100g",
        "beta_carotene_100g",
        "vitamin_b1_100g",
        "vitamin_b2_100g",
        "vitamin_b3_100g",
        "vitamin_b6_100g",
        "vitamin_b9_100g",
        "vitamin_b12_100g",
        "pantothenic_acid_100g",
        "vitamin_c_100g",
        "vitamin_d_100g",
        "vitamin_e_100g",
        "potassium_100g",
        "sodium_100g",
        "chloride_100g",
        "calcium_100g",
        "magnesium_100g",
        "phosphorus_100g",
        "iron_100g",
        "iodine_100g",
        "zinc_100g",
        "selenium_100g",
    ]

    # Multipliers convert the source unit to grams per 100g for downstream parity.
    UNIT_MULTIPLIERS: dict[str, Decimal] = {
        # milligrams -> grams
        "cholesterol_100g": Decimal("0.001"),
        "vitamin_b1_100g": Decimal("0.001"),
        "vitamin_b2_100g": Decimal("0.001"),
        "vitamin_b3_100g": Decimal("0.001"),
        "vitamin_b6_100g": Decimal("0.001"),
        "pantothenic_acid_100g": Decimal("0.001"),
        "vitamin_c_100g": Decimal("0.001"),
        "vitamin_e_100g": Decimal("0.001"),
        "potassium_100g": Decimal("0.001"),
        "sodium_100g": Decimal("0.001"),
        "chloride_100g": Decimal("0.001"),
        "calcium_100g": Decimal("0.001"),
        "magnesium_100g": Decimal("0.001"),
        "phosphorus_100g": Decimal("0.001"),
        "iron_100g": Decimal("0.001"),
        "zinc_100g": Decimal("0.001"),
        # micrograms -> grams
        "vitamin_a_100g": Decimal("0.000001"),
        "beta_carotene_100g": Decimal("0.000001"),
        "vitamin_b9_100g": Decimal("0.000001"),
        "vitamin_b12_100g": Decimal("0.000001"),
        "vitamin_d_100g": Decimal("0.000001"),
        "iodine_100g": Decimal("0.000001"),
        "selenium_100g": Decimal("0.000001"),
    }

    QUANTIZE_TARGET = Decimal("0.000000001")

    def __init__(self, source_path: Path, target_path: Path):
        """
        Initializes the converter with the source Excel file and target CSV file paths.

        :param source_path: Path to the source Excel file.
        :type source_path: Path
        :param target_path: Path to the target CSV file.
        :type target_path: Path
        """

        self._source_path = source_path
        self._target_path = target_path

    def convert(
        self,
        progress_callback: ProgressCallback | None = None,
    ) -> None:
        """
        Converts the Swiss nutrition Excel file to a CSV file with the standardized output structure.

        :raises OSError: If the source file cannot be read or the target file cannot be written.
        :raises ValueError: If the expected header row cannot be found or required columns are missing
        """

        workbook = load_workbook(self._source_path, data_only=True, read_only=True)
        worksheet = workbook.active

        if worksheet is None:
            raise ValueError("The Excel file does not contain any worksheets.")

        rows = list(worksheet.iter_rows(values_only=True))
        total_size = len(rows)

        emit_progress(
            progress_callback,
            phase="conversion",
            description=f"Converting {self._source_path} to CSV",
            completed=0,
            total=total_size,
        )

        header_index = self._find_header_index(rows)
        header_row = rows[header_index]

        sanitized_headers = [
            self._sanitize_header(str(cell)) if cell is not None else ""
            for cell in header_row
        ]

        column_indexes: dict[str, int] = {}
        self._map_columns_to_output_structure(sanitized_headers, column_indexes)

        with self._target_path.open("w", newline="", encoding="utf-8") as csv_file:
            writer = csv.DictWriter(csv_file, fieldnames=self.OUTPUT_COLUMNS)
            writer.writeheader()

            for row_index, row in enumerate(rows[header_index + 1 :]):
                if row is None:
                    continue

                mapped_row = {}
                empty_row = True

                for column_name, source_index in column_indexes.items():
                    value = row[source_index] if source_index < len(row) else None
                    text_value = ""

                    if value is not None:
                        text_value = str(value).strip()
                        normalized_value = text_value.lower().replace(" ", "")

                        if normalized_value in {"k.a.", "k.a"}:
                            text_value = ""
                        elif normalized_value in {"sp.", "sp"}:
                            text_value = ""
                        else:
                            text_value = (
                                text_value.replace("<", "").replace("≤", "").strip()
                            )
                            text_value = self._normalize_unit(column_name, text_value)

                    if text_value:
                        empty_row = False

                    mapped_row[column_name] = text_value

                if empty_row:
                    continue

                writer.writerow(
                    {
                        column: mapped_row.get(column, "")
                        for column in self.OUTPUT_COLUMNS
                    }
                )

                emit_progress(
                    progress_callback,
                    phase="conversion",
                    description=f"Converting {self._source_path} to CSV",
                    completed=row_index,
                    total=total_size,
                )

            emit_progress(
                progress_callback,
                phase="conversion",
                description=f"Converting {self._source_path} to CSV",
                completed=total_size,
                total=total_size,
            )

    def _find_header_index(self, rows):
        """
        Finds the index of the header row in the given rows. The header row is identified by the presence
        of a cell in the first column that contains "ID" (case-insensitive).

        :param rows: The rows to search through.
        :type rows: list of tuples representing Excel rows
        :return: The index of the header row.
        :rtype: int
        :raises ValueError: If no header row with an "ID" column is found.
        """
        header_index = None

        for idx, row in enumerate(rows):
            if not row:
                continue

            first_cell = row[0]

            if first_cell and str(first_cell).strip().lower() == "id":
                header_index = idx
                break

        if header_index is None:
            raise ValueError(
                "Could not locate header row with an 'ID' column in the Excel file."
            )

        return header_index

    def _sanitize_header(self, label: str) -> str:
        """
        Sanitizes a header label by replacing special characters with ASCII equivalents,
        converting to lowercase, and replacing non-alphanumeric characters with underscores.

        :param label: The header label to sanitize.
        :type label: str
        :return: The sanitized header label.
        :rtype: str
        """
        replacements = {
            "ä": "ae",
            "Ä": "Ae",
            "ö": "oe",
            "Ö": "Oe",
            "ü": "ue",
            "Ü": "Ue",
            "ß": "ss",
        }

        processed = label

        for original, replacement in replacements.items():
            processed = processed.replace(original, replacement)

        normalized = unicodedata.normalize("NFKD", processed)
        ascii_only = normalized.encode("ascii", "ignore").decode("ascii")
        lowercased = ascii_only.lower()
        replaced = lowercased.replace("%", "percent")
        collapsed = re.sub(r"[^a-z0-9]+", "_", replaced)
        return collapsed.strip("_")

    def _normalize_unit(self, column_name: str, text_value: str) -> str:
        """
        Converts numeric text in alternate units to grams where required.

        :param column_name: The name of the column being processed.
        :type column_name: str
        :param text_value: The numeric text value to convert.
        :type text_value: str
        :return: The converted value as a string.
        :rtype: str
        """

        multiplier = self.UNIT_MULTIPLIERS.get(column_name)
        if multiplier is None or text_value == "":
            return text_value

        normalized = text_value.replace(" ", "").replace(",", ".").replace("'", "")

        try:
            number = Decimal(normalized)
        except InvalidOperation:
            return text_value

        converted = number * multiplier
        if converted == 0:
            return "0"

        quantized = converted.quantize(self.QUANTIZE_TARGET, rounding=ROUND_HALF_UP)
        formatted = f"{quantized:.9f}".rstrip("0").rstrip(".")
        return formatted if formatted else "0"

    def _map_columns_to_output_structure(self, sanitized_headers, column_indexes):
        """
        Maps sanitized header labels to the output structure defined by COLUMN_MAP and OUTPUT_COLUMNS.

        :param sanitized_headers: The list of sanitized header labels.
        :type sanitized_headers: list of str
        :param column_indexes: The dictionary to populate with column indexes.
        :type column_indexes: dict
        :raises ValueError: If any required columns defined in OUTPUT_COLUMNS are missing from the input.
        """
        for index, sanitized in enumerate(sanitized_headers):
            if sanitized in self.COLUMN_MAP:
                column_name = self.COLUMN_MAP[sanitized]

                if column_name not in column_indexes:
                    column_indexes[column_name] = index

        missing_columns = [
            column for column in self.OUTPUT_COLUMNS if column not in column_indexes
        ]

        if missing_columns:
            raise ValueError(
                f"Missing columns in Swiss nutrition export: {', '.join(missing_columns)}"
            )
